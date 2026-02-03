#!/usr/bin/env python3
"""
엑셀 파일 시트 병합 프로그램
- 여러 엑셀 파일의 첫 번째 시트를 하나의 엑셀 파일로 병합
- 파일명에서 구간명 추출하여 시트명으로 사용
- 구간별 O열(DIA/INCH) 합산 요약 시트 생성
"""

import os
import sys
import re
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from datetime import datetime


def extract_section_name(filename):
    """파일명에서 구간명 추출 (예: 12345_구간명.xlsx -> 구간명)"""
    base_name = os.path.splitext(os.path.basename(filename))[0]

    # '_' 뒤의 구간명 추출
    if '_' in base_name:
        section = base_name.split('_', 1)[1]
        return section
    return base_name


def is_valid_data_sheet(sheet):
    """유효한 데이터 시트인지 확인 (A1: 등록여부, B1: SBNM 등)"""
    try:
        a1 = sheet['A1'].value
        b1 = sheet['B1'].value

        # A1에 '등록여부', B1에 'SBNM' 등이 있으면 유효한 시트
        if a1 and '등록' in str(a1):
            return True
        if b1 and 'SBNM' in str(b1).upper():
            return True
        # 데이터가 있으면 일단 유효하다고 판단
        if a1 is not None or b1 is not None:
            return True
    except:
        pass
    return False


def get_first_valid_sheet(workbook):
    """워크북에서 첫 번째 유효한 데이터 시트 반환"""
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        if is_valid_data_sheet(sheet):
            return sheet
    # 유효한 시트가 없으면 첫 번째 시트 반환
    return workbook.worksheets[0]


def calculate_o_column_sum(sheet):
    """O열(15번째 열)의 DIA/INCH 합산 계산"""
    total = 0
    o_col = 15  # O열은 15번째

    for row in range(2, sheet.max_row + 1):  # 헤더 제외, 2행부터
        cell_value = sheet.cell(row=row, column=o_col).value
        if cell_value is not None:
            try:
                # 숫자로 변환 시도
                num_value = float(str(cell_value).replace(',', ''))
                total += num_value
            except (ValueError, TypeError):
                pass

    return total


def copy_sheet_data(source_sheet, target_sheet):
    """소스 시트의 데이터를 타겟 시트로 복사"""
    for row_idx, row in enumerate(source_sheet.iter_rows(), 1):
        for col_idx, cell in enumerate(row, 1):
            target_cell = target_sheet.cell(row=row_idx, column=col_idx)
            target_cell.value = cell.value

            # 스타일 복사 시도
            if cell.has_style:
                try:
                    target_cell.font = Font(
                        name=cell.font.name,
                        size=cell.font.size,
                        bold=cell.font.bold,
                        italic=cell.font.italic,
                        color=cell.font.color
                    )
                    target_cell.alignment = Alignment(
                        horizontal=cell.alignment.horizontal,
                        vertical=cell.alignment.vertical,
                        wrap_text=cell.alignment.wrap_text
                    )
                except:
                    pass

    # 열 너비 복사
    for col_idx, col_dim in source_sheet.column_dimensions.items():
        if col_dim.width:
            target_sheet.column_dimensions[col_idx].width = col_dim.width


def create_summary_sheet(workbook, section_sums):
    """구간별 DIA/INCH 합산 요약 시트 생성"""
    summary = workbook.create_sheet("요약", 0)  # 맨 앞에 생성

    # 스타일 정의
    header_font = Font(bold=True, size=12)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font_white = Font(bold=True, size=12, color="FFFFFF")
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    center_align = Alignment(horizontal='center', vertical='center')

    # 제목
    summary['A1'] = "구간별 DIA/INCH 합산 요약"
    summary['A1'].font = Font(bold=True, size=14)
    summary.merge_cells('A1:C1')

    # 생성 일시
    summary['A2'] = f"생성일시: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    summary.merge_cells('A2:C2')

    # 헤더
    headers = ['No.', '구간명', 'DIA/INCH 합계']
    for col, header in enumerate(headers, 1):
        cell = summary.cell(row=4, column=col, value=header)
        cell.font = header_font_white
        cell.fill = header_fill
        cell.border = border
        cell.alignment = center_align

    # 데이터
    total_sum = 0
    for idx, (section, sum_value) in enumerate(section_sums.items(), 1):
        row = idx + 4

        # No.
        cell = summary.cell(row=row, column=1, value=idx)
        cell.border = border
        cell.alignment = center_align

        # 구간명
        cell = summary.cell(row=row, column=2, value=section)
        cell.border = border
        cell.alignment = center_align

        # 합계
        cell = summary.cell(row=row, column=3, value=sum_value)
        cell.border = border
        cell.alignment = center_align
        cell.number_format = '#,##0.00'

        total_sum += sum_value

    # 총합계
    total_row = len(section_sums) + 5
    summary.cell(row=total_row, column=1, value="").border = border

    total_cell = summary.cell(row=total_row, column=2, value="총 합계")
    total_cell.font = Font(bold=True)
    total_cell.border = border
    total_cell.alignment = center_align

    total_value_cell = summary.cell(row=total_row, column=3, value=total_sum)
    total_value_cell.font = Font(bold=True)
    total_value_cell.border = border
    total_value_cell.alignment = center_align
    total_value_cell.number_format = '#,##0.00'

    # 열 너비 조정
    summary.column_dimensions['A'].width = 8
    summary.column_dimensions['B'].width = 30
    summary.column_dimensions['C'].width = 18

    return summary


def merge_excel_files(input_files, output_file="merged_output.xlsx"):
    """
    여러 엑셀 파일을 하나로 병합

    Args:
        input_files: 입력 엑셀 파일 경로 리스트
        output_file: 출력 엑셀 파일 경로
    """
    if not input_files:
        print("오류: 입력 파일이 없습니다.")
        return False

    # 새 워크북 생성
    merged_wb = Workbook()
    # 기본 시트 삭제
    merged_wb.remove(merged_wb.active)

    section_sums = {}  # 구간별 합산 저장

    print(f"\n{'='*50}")
    print("엑셀 파일 병합 시작")
    print(f"{'='*50}\n")

    for idx, file_path in enumerate(input_files, 1):
        if not os.path.exists(file_path):
            print(f"경고: 파일을 찾을 수 없습니다 - {file_path}")
            continue

        try:
            print(f"[{idx}/{len(input_files)}] 처리 중: {os.path.basename(file_path)}")

            # 파일 로드
            source_wb = load_workbook(file_path, data_only=True)

            # 구간명 추출
            section_name = extract_section_name(file_path)

            # 시트명이 중복될 경우 처리
            original_section = section_name
            counter = 1
            while section_name in [s.title for s in merged_wb.worksheets]:
                section_name = f"{original_section}_{counter}"
                counter += 1

            # 첫 번째 유효한 시트 가져오기
            source_sheet = get_first_valid_sheet(source_wb)

            # 새 시트 생성 및 데이터 복사
            target_sheet = merged_wb.create_sheet(title=section_name[:31])  # 시트명 31자 제한
            copy_sheet_data(source_sheet, target_sheet)

            # O열 합계 계산
            o_sum = calculate_o_column_sum(source_sheet)
            section_sums[section_name] = o_sum

            print(f"  → 구간명: {section_name}")
            print(f"  → O열 합계: {o_sum:,.2f}")
            print()

            source_wb.close()

        except Exception as e:
            print(f"오류: {file_path} 처리 중 문제 발생 - {str(e)}")
            continue

    if not merged_wb.worksheets:
        print("오류: 병합할 시트가 없습니다.")
        return False

    # 요약 시트 생성 (맨 앞에)
    create_summary_sheet(merged_wb, section_sums)

    # 저장
    merged_wb.save(output_file)

    print(f"{'='*50}")
    print(f"병합 완료!")
    print(f"출력 파일: {output_file}")
    print(f"총 {len(section_sums)}개 구간 처리됨")
    print(f"{'='*50}\n")

    return True


def main():
    """메인 함수"""
    print("\n" + "="*50)
    print("  엑셀 파일 시트 병합 프로그램")
    print("="*50)

    # 명령줄 인수로 파일 받기
    if len(sys.argv) > 1:
        input_files = sys.argv[1:]
        output_file = "merged_output.xlsx"
    else:
        # 대화형 모드
        print("\n병합할 엑셀 파일들을 입력하세요.")
        print("(파일 경로를 한 줄에 하나씩 입력, 빈 줄 입력시 완료)")
        print("또는 현재 폴더의 모든 xlsx 파일을 병합하려면 'all' 입력\n")

        input_files = []
        while True:
            file_input = input("파일 경로: ").strip()
            if not file_input:
                break
            if file_input.lower() == 'all':
                # 현재 폴더의 모든 xlsx 파일
                input_files = [f for f in os.listdir('.')
                              if f.endswith('.xlsx') and not f.startswith('merged_')]
                input_files.sort()
                break
            if os.path.exists(file_input):
                input_files.append(file_input)
            else:
                print(f"  → 파일을 찾을 수 없습니다: {file_input}")

        # 출력 파일명 입력
        output_input = input("\n출력 파일명 (기본값: merged_output.xlsx): ").strip()
        output_file = output_input if output_input else "merged_output.xlsx"
        if not output_file.endswith('.xlsx'):
            output_file += '.xlsx'

    if not input_files:
        print("\n오류: 입력 파일이 없습니다.")
        sys.exit(1)

    print(f"\n입력 파일 목록:")
    for f in input_files:
        print(f"  - {f}")

    # 병합 실행
    success = merge_excel_files(input_files, output_file)

    if success:
        print("프로그램이 정상적으로 완료되었습니다.")
    else:
        print("프로그램 실행 중 오류가 발생했습니다.")
        sys.exit(1)


if __name__ == "__main__":
    main()
